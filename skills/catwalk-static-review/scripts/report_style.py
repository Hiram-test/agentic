"""Measured original-report typography: SimSun / SimHei / Times New Roman."""
from pathlib import Path
from datetime import datetime,timezone
import re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER,TA_JUSTIFY
from reportlab.pdfgen import canvas
from reportlab.platypus import BaseDocTemplate,PageTemplate,Frame,Paragraph,Spacer,PageBreak,Table,TableStyle,Image,KeepTogether

FONT_ROOT=Path(__file__).resolve().parents[1]/'assets/fonts'
W,H=595.32,842.04
LEFT,RIGHT,TOP,BOTTOM=70.87,42.52,63.0,83.0
WIDTH=W-LEFT-RIGHT

def register_fonts():
    for name,file in [('ReportSong','simsun.ttf'),('ReportHei','simhei.ttf'),('ReportLatin','times.ttf'),('ReportLatinBold','timesbd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(FONT_ROOT/file)))

def rich(text,bold=False):
    # Original uses Times New Roman for Latin/numbers; preserve the Chinese style font.
    result=[]
    for token in re.split(r'([⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺]+|[\x20-\x7e\u0370-\u03ff]+)',str(text)):
        if not token:continue
        if all(c in '⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺' for c in token):
            result.append('<super><font name="ReportLatin">'+token.translate(str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺','0123456789-+'))+'</font></super>');continue
        if token.isascii() or all(ord(c)<128 or '\u0370'<=c<='\u03ff' for c in token):result.append('<font name="'+('ReportLatinBold' if bold and all(c in ' 0123456789.-≥' for c in token) else 'ReportLatin')+'">'+escape(token)+'</font>')
        else:result.append(escape(token))
    return ''.join(result).replace('\n','<br/>')

class NumberedCanvas(canvas.Canvas):
    def __init__(self,*a,**kw):super().__init__(*a,**kw);self.saved=[]
    def showPage(self):self.saved.append(dict(self.__dict__));self._startPage()
    def save(self):
        total=len(self.saved)
        for state in self.saved:
            self.__dict__.update(state);self.setFont('ReportLatin',9)
            self.drawRightString(W-RIGHT,61,f'{self._pageNumber:03d}/{total:03d}')
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

class PDF(BaseDocTemplate):
    def __init__(self,path):
        super().__init__(str(path),pagesize=(W,H),leftMargin=LEFT,rightMargin=RIGHT,topMargin=TOP,bottomMargin=BOTTOM)
        self.addPageTemplates(PageTemplate(id='normal',frames=Frame(LEFT,BOTTOM,WIDTH,H-TOP-BOTTOM,id='body',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0),onPage=self.decorate))
    def decorate(self,c,d):
        if d.page>1:
            c.setStrokeColor(colors.black);c.setLineWidth(.5);c.line(69.5,H-56.1,554.3,H-56.1);c.line(69.5,H-768.7,554.3,H-768.7)
            c.setFont('ReportSong',8);c.drawString(LEFT,61,'猫道及门架承重索静力复核 / P1—P6 六工况')
    def afterFlowable(self,f):
        if isinstance(f,Paragraph) and f.style.name in ('H1','H2'):
            self.notify('TOCEntry',(0 if f.style.name=='H1' else 1,f.getPlainText(),self.page))

def setfont(run,cn='宋体',size=12,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cn)

def field(par,instruction):
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),instruction);par._p.append(f)

class Writer:
    def __init__(self,out):
        register_fonts();self.out=Path(out);self.doc=Document();self.story=[];self.chapter=1;self.table_no=0;self.fig_no=0;self.section=''
        self.normal=ParagraphStyle('Body',fontName='ReportSong',fontSize=12,leading=23.4,firstLineIndent=24,spaceAfter=0,wordWrap='CJK',alignment=TA_JUSTIFY)
        self.cell=ParagraphStyle('Cell',fontName='ReportSong',fontSize=10.5,leading=14,spaceAfter=0,alignment=TA_CENTER,wordWrap='CJK')
        self.caption=ParagraphStyle('Caption',parent=self.cell,leading=18,spaceBefore=4,spaceAfter=6,keepWithNext=True)
        self.head={1:ParagraphStyle('H1',fontName='ReportHei',fontSize=16,leading=24,alignment=TA_CENTER,spaceBefore=20,spaceAfter=22,keepWithNext=True),2:ParagraphStyle('H2',fontName='ReportHei',fontSize=14,leading=23.4,spaceBefore=12,spaceAfter=7,keepWithNext=True),3:ParagraphStyle('H3',fontName='ReportHei',fontSize=12,leading=23.4,spaceBefore=7,spaceAfter=4,keepWithNext=True)}
        sec=self.doc.sections[0];sec.page_width=Pt(W);sec.page_height=Pt(H);sec.top_margin=Pt(TOP);sec.bottom_margin=Pt(BOTTOM);sec.left_margin=Pt(LEFT);sec.right_margin=Pt(RIGHT);sec.header_distance=Pt(48);sec.footer_distance=Pt(57);sec.different_first_page_header_footer=True
        for name in ['Normal','Heading 1','Heading 2','Heading 3']:
            sty=self.doc.styles[name];sty.font.name='Times New Roman';sty._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体');sty.font.size=Pt(12 if name in ('Normal','Heading 3') else 16 if name=='Heading 1' else 14)
            sty.font.bold=False;pf=sty.paragraph_format;pf.line_spacing=Pt(23.4);pf.space_after=Pt(0);pf.first_line_indent=Pt(24 if name=='Normal' else 0)
        for name in ['Heading 1','Heading 2','Heading 3']:
            pf=self.doc.styles[name].paragraph_format;pf.space_before=Pt(20 if name=='Heading 1' else 10);pf.space_after=Pt(22 if name=='Heading 1' else 6)
            if name=='Heading 1':pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for par,border in [(sec.header.paragraphs[0],'bottom'),(sec.footer.paragraphs[0],'top')]:
            pb=OxmlElement('w:pBdr');b=OxmlElement('w:'+border)
            for k,v in {'val':'single','sz':'4','space':'4','color':'000000'}.items():b.set(qn('w:'+k),v)
            pb.append(b);par._p.get_or_add_pPr().append(pb);par.paragraph_format.first_line_indent=Pt(0)
        footer=sec.footer.paragraphs[0];setfont(footer.add_run('猫道及门架承重索静力复核 / P1—P6 六工况'),size=8)
        footer.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH-60));footer.add_run('\t');field(footer,'PAGE \\# "000"');footer.add_run('/');field(footer,'NUMPAGES \\# "000"')
        first=sec.first_page_footer.paragraphs[0];first.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field(first,'PAGE \\# "000"');first.add_run('/');field(first,'NUMPAGES \\# "000"')
    def cover(self,run):
        self.story.append(Spacer(1,68))
        for text,size,gap in [('张靖皋长江大桥南航道桥',42,24),('猫道结构复核计算报告',36,32)]:
            style=ParagraphStyle('Cover',fontName='ReportHei',fontSize=size,leading=size+8,alignment=TA_CENTER)
            self.story.extend([Paragraph(text,style),Spacer(1,gap)])
        self.story.append(Paragraph('P1—P6 六工况静力复算',ParagraphStyle('Subtitle',fontName='ReportSong',fontSize=14,leading=20,alignment=TA_CENTER)))
        self.story.append(Spacer(1,372))
        self.story.append(Paragraph('猫道及门架承重索整体计算',ParagraphStyle('Organisation',fontName='ReportHei',fontSize=14,leading=22,alignment=TA_CENTER)))
        self.story.append(Spacer(1,20));date=datetime.now(timezone.utc);cn='○一二三四五六七八九';year=''.join(cn[int(x)] for x in str(date.year));month=['','一','二','三','四','五','六','七','八','九','十','十一','十二'][date.month];day=(cn[date.day] if date.day<10 else '十' if date.day==10 else ('十'+cn[date.day-10] if date.day<20 else '二十'+(cn[date.day-20] if date.day>20 else ''))) if date.day<30 else '三十'+(cn[date.day-30] if date.day>30 else '')
        label=f'{year}年{month}月{day}日'
        self.story.append(Paragraph(label,ParagraphStyle('Date',fontName='ReportHei',fontSize=14,leading=22,alignment=TA_CENTER)))
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(50);p.paragraph_format.space_before=Pt(68);p.paragraph_format.space_after=Pt(24);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;setfont(p.add_run('张靖皋长江大桥南航道桥'),'黑体',42)
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(44);p.paragraph_format.space_after=Pt(32);p.alignment=1;setfont(p.add_run('猫道结构复核计算报告'),'黑体',36)
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=1;setfont(p.add_run('P1—P6 六工况静力复算'),size=14)
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(372);p.alignment=1;setfont(p.add_run('猫道及门架承重索整体计算'),'黑体',14)
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(20);p.alignment=1;setfont(p.add_run(label),'黑体',14)
    def title(self,text):
        p=self.doc.add_paragraph();p.alignment=1;p.paragraph_format.first_line_indent=Pt(0);setfont(p.add_run(text),'黑体',15)
        self.story.append(Paragraph(text,ParagraphStyle('Title',fontName='ReportHei',fontSize=15,leading=24,alignment=TA_CENTER,spaceAfter=14)))
    def p(self,text):
        par=self.doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;setfont(par.add_run(str(text)))
        self.story.append(Paragraph(rich(text),self.normal))
    def h(self,text,level=1):
        if level==1:
            match=re.match(r'第\s*(\d+)\s*章',text)
            if match:self.chapter=int(match[1]);self.table_no=0;self.fig_no=0
        self.section=re.sub(r'^[\d.]+\s*','',text)
        self.doc.add_heading(text,level);self.story.append(Paragraph(rich(text),self.head[level]))
    def caption_p(self,text):
        par=self.doc.add_paragraph();par.alignment=1;par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.keep_with_next=True;setfont(par.add_run(text),size=10.5)
        self.story.append(Paragraph(rich(text),self.caption))
    def table(self,headers,rows,widths=None,caption=None,spans=None,table_width=None):
        self.table_no+=1;self.caption_p(f'表{self.chapter}-{self.table_no}  '+(caption or self.section))
        n=len(headers);widths=widths or ([WIDTH*.30,WIDTH*.70] if n==2 else [WIDTH/n]*n);widths=[x*(table_width or WIDTH)/sum(widths) for x in widths]
        t=self.doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Pt(width)
        for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
        header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
        for row in rows:
            for c,s in zip(t.add_row().cells,row):c.text=str(s)
        for row in t.rows:
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.alignment=1;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(14);p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                    for r in p.runs:setfont(r,size=10.5)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            item=OxmlElement('w:'+side)
            for k,v in {'val':'single','sz':'4','color':'666666'}.items():item.set(qn('w:'+k),v)
            borders.append(item)
        t._tbl.tblPr.append(borders)
        data=[[Paragraph(rich(x,bold=len(row)>1 and row[1]=='安全系数'),self.cell) for x in row] for row in [headers]+rows]
        for index,row in enumerate(rows,1):
            if len(row)>1 and row[1]=='安全系数':
                for c in t.rows[index].cells[2:]:
                    for p in c.paragraphs:
                        for run in p.runs:run.bold=True
        tab=Table(data,colWidths=widths,repeatRows=1,hAlign='CENTER')
        cmds=[('GRID',(0,0),(-1,-1),.4,colors.HexColor('#666666')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),2.5),('BOTTOMPADDING',(0,0),(-1,-1),2.5)]
        for start,end in spans or []:
            cmds.append(('SPAN',start,end));t.cell(start[1],start[0]).merge(t.cell(end[1],end[0]))
        tab.setStyle(TableStyle(cmds));self.story.extend([tab,Spacer(1,10)])
    def strength_table(self,rows,typ):
        body=[];spans=[]
        for i,r in enumerate(rows):
            values=[s for s in r['span_forces'] if s['type']==typ]
            body.extend([[f'工况{i+1}','承重索轴力MaxN（kN）',*[f"{s['N_kN']:.2f}" for s in values]],['','承重索破断力（kN）',*[f"{s['breaking_force_kN']:.0f}" for s in values]],['','安全系数',*[f"{s['safety_factor']:.3f}" for s in values]]]);spans.append(((0,1+3*i),(0,3+3*i)))
        name='猫道承重索' if typ=='bottom' else '门架承重索'
        self.table(['','', '北边跨','主跨','南边跨','南辅跨'],body,widths=[54,139,54,54,54,54],caption=name+'强度验算表',spans=spans,table_width=409)
        self.table(['工况','北边跨','主跨','南边跨','南辅跨'],[[r['case'],*[f"{s['error_percent']:+.3f}%" for s in r['span_forces'] if s['type']==typ]] for r in rows],caption=name+'轴力与原报告对照（相对差）')
        for r in rows:
            vals=[s for s in r['span_forces'] if s['type']==typ];k=min(x['safety_factor'] for x in vals);req=vals[0]['required_factor']
            self.p(f"工况{r['case'][1:]}：最小安全系数{k:.3f}，要求不小于{req:.1f}，"+('满足要求。' if k>=req else '不满足要求。'))
    def image(self,path,caption):
        self.fig_no+=1;cap=f'图{self.chapter}-{self.fig_no}  {caption}'
        p=self.doc.add_paragraph();p.alignment=1;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Pt(WIDTH))
        p=self.doc.add_paragraph();p.alignment=1;p.paragraph_format.first_line_indent=Pt(0);setfont(p.add_run(cap),size=10.5)
        style=ParagraphStyle('FigureCaption',parent=self.caption,keepWithNext=False,spaceAfter=10)
        self.story.append(KeepTogether([Image(str(path),width=WIDTH,height=WIDTH*3.8/12),Paragraph(rich(cap),style)]))
    def page(self):self.doc.add_page_break();self.story.append(PageBreak())
    def embed_word_fonts(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree
        import uuid
        fontpart=self.doc.part.part_related_by(RT.FONT_TABLE)
        tree=etree.fromstring(fontpart.blob)
        for family,filename,bold in [('宋体','simsun.ttf',False),('黑体','simhei.ttf',False),('Times New Roman','times.ttf',False),('Times New Roman','timesbd.ttf',True)]:
            matches=tree.findall(qn('w:font'))
            entry=next((x for x in matches if x.get(qn('w:name'))==family),None)
            if entry is None:entry=OxmlElement('w:font');entry.set(qn('w:name'),family);tree.append(entry)
            key=uuid.uuid4();data=bytearray((FONT_ROOT/filename).read_bytes());mask=key.bytes[::-1]
            for i in range(32):data[i]^=mask[i%16]
            part=Part(PackURI('/word/fonts/'+filename+'.odttf'),'application/vnd.openxmlformats-officedocument.obfuscatedFont',bytes(data),self.doc.part.package)
            rid=fontpart.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font')
            item=OxmlElement('w:embedBold' if bold else 'w:embedRegular');item.set(qn('r:id'),rid);item.set(qn('w:fontKey'),'{'+str(key).upper()+'}');entry.append(item)
        fontpart._blob=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
        self.doc.settings.element.append(OxmlElement('w:embedTrueTypeFonts'))
    def save(self):
        self.embed_word_fonts()
        self.doc.save(self.out/'catwalk_static_review.docx');PDF(self.out/'catwalk_static_review.pdf').multiBuild(self.story,canvasmaker=NumberedCanvas)
