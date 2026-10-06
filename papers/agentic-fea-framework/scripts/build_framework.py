"""Build the reference-led manuscript framework and figure study from checked evidence."""
from pathlib import Path
import csv, json, hashlib
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
SKILL=ROOT.parents[1]/'skills/catwalk-static-review'
SUM=json.loads((SKILL/'examples/summary.json').read_text())
RUN=json.loads((SKILL/'examples/run.json').read_text())
OUT=ROOT/'figures'; OUT.mkdir(exist_ok=True)
REF='Qi, Y., Xu, R., and Chu, X. FeaGPT: an End-to-End agentic-AI for Finite Element Analysis. arXiv:2510.21993v1 (2025). https://arxiv.org/abs/2510.21993v1'
FIGS=[
(1,4,'Workflow architecture','A user request feeds the chief engineer/planner; a knowledge base supplies geometry, materials and prior cases; four connected modules perform geometry, meshing, simulation and data analysis. Arrows show both task flow and returned results.',
 'Keep the left-to-right request-to-result composition and the knowledge block above the workflow. Replace the aircraft request with the catwalk INP set and reference report. Use four module boxes: input/model audit; Gmsh topology reconstruction; nonlinear CalculiX execution; field recovery and report generation. Add the manifest, Skill and benchmark assets above, and the Word/PDF/CSV/VTU outputs at the return arrow.',
 'SKILL.md; scripts/run.py; assets/reference/manifest.json; examples/run.json',
 'Figure 1. Skill-guided workflow for nonlinear static verification of the catwalk cable system, connecting versioned inputs, mesh verification, finite element solution, field recovery and engineering-report generation.'),
(2,12,'Geometry and structural detail','Three panels progress from an external wing view with a highlighted detail to a complete assembly and an internal structural view. The red ellipse and arrow connect global geometry to a local feature.',
 'Retain three panels and the global-to-local zoom. Panel (a): full X–Z catwalk profile and control points. Panel (b): bottom cables, portal cables and equivalent members in three colors. Panel (c): enlarged portal/cable connection with original node and element IDs, material assignments and supports. Draw geometry directly from the INP coordinates.',
 'assets/inputs/migrate_P1.inp; scripts/model.py; assets/reference/P1_input_audit.json',
 'Figure 2. Catwalk finite element model: (a) longitudinal profile and control points; (b) cable families and equivalent members; (c) a portal connection with element assignments and boundary conditions.'),
(3,13,'Mesh-to-field correspondence','A mesh overview and local mesh enlargement occupy the left side; a stress contour occupies the right. The arrangement connects discrete geometry to the solved response.',
 'Retain the two main panels and inset. Panel (a): original/Gmsh line mesh overlay with an enlarged tag-mapping detail. Annotate 1,125 nodes, 1,194 elements and zero coordinate discrepancy. Panel (b): P5 axial-force line contour using the same geometry; show kN, span group and governing element. Draw only line elements.',
 'examples/validated-run.zip: mesh/mesh_map.json and P5 field files; scripts/mesh_gmsh.py; scripts/postprocess.py',
 'Figure 3. Mesh reconstruction and field transfer: (a) Gmsh mesh and original-input tag correspondence; (b) P5 bottom-cable axial-force distribution on the verified line mesh.'),
(4,15,'Four-panel quantitative evidence','The original 2×2 panel layout moves through parameter correlations, a stress–weight Pareto scatter, load-step stress distributions and fatigue-life results. Each panel answers a different quantitative question.',
 'Keep the 2×2 evidence layout. Panel (a): six-case CalculiX/MAPDL peak-displacement bars. Panel (b): parity scatter for 48 report/calculated span-force pairs, distinguished by cable family. Panel (c): minimum four-span safety factor by family and case with required values. Panel (d): signed relative-error heatmap for six cases × eight span groups. Use the actual summary JSON for every point.',
 'examples/summary.json; examples/comparison.csv; figures/figure_4_data.csv',
 'Figure 4. Six-case verification: (a) maximum displacement compared with MAPDL; (b) 48 span-force comparisons with the engineering report; (c) minimum span safety factors and required values; (d) signed force differences by case and span.'),
(5,18,'Grouped field visualization','Six panels are arranged in three columns and two rows. Compressor and turbine results combine full-component/sector views and stress/displacement fields with independent color bars.',
 'Keep six panels, arranged as three columns × two rows. Columns: P1 dead load, P5 maximum gust and P6 cooling. Upper row: signed UZ. Lower row: bottom-cable axial force. Use common limits within each field row, label mm/kN, and annotate the relevant extreme node/element. Place all six load cases in supplementary field plates using the same settings.',
 'examples/validated-run.zip: P1/P5/P6 nodes.csv, element results and VTU; scripts/postprocess.py',
 'Figure 5. Representative field responses under dead load, maximum gust and cooling: upper row, vertical displacement; lower row, bottom-cable axial force. Columns correspond to P1, P5 and P6.')]
SECTIONS=[
('I. Introduction','I. Introduction, pp. 1–2',[
 'Paragraph 1 — Introduce agentic engineering workflows through the sequence used in FeaGPT: language-based task specification, tool execution and computational engineering results. Cite FeaGPT at the transition to finite element analysis.',
 'Paragraph 2 — Move from the reference aircraft/component setting to bridge catwalk verification. Describe the practical input chain: a prestressed INP model, six load combinations, an existing engineering report and a required reproducible calculation package.',
 'Paragraph 3 — Define the central workflow problem: preserving model identity, prestress, stepwise total loads and result-group definitions across automated execution. Use the P5 file-version/load-history episode as the motivating example.',
 'Paragraph 4 — Introduce the Skill as the executable analysis specification, then list the paper contributions in the reference order: structured task execution; Gmsh–CalculiX integration; six-case quantitative verification; automatic engineering-report delivery.'
]),
('II.A. System overview','II.A. System Overview, p. 3; Fig. 1',[
 'Describe the four workflow modules and their file interfaces using Figure 1. The assistant reads the English Skill and invokes the packaged tools; each tool produces named, machine-readable evidence.',
 'Identify the model asset manifest, original PDF, benchmark JSON and font/style assets. Trace one request from asset verification to the report and raw solver package.'
]),
('II.B. Skill-guided planning and task orchestration','II.B. Engineering Analysis Planning and Task Orchestration, pp. 3–6; Algorithm 1',[
 'Replace the reference geometry-design JSON with the catwalk task schema: input manifest, cases P1–P6, solver identity, units, execution stages, comparison criteria and report outputs.',
 'Present Algorithm 1 as the executable order: validate → reconstruct mesh → solve each case → read final fields → recover forces → compare → generate report → hash outputs. State the stage-specific completion conditions.',
 'For the planned GLM execution experiment, record provider/model identifier, request, tool calls, elapsed time, retries and final artifact status in a session log. Fill the model identifier from the actual run.'
]),
('II.C. Versioned model reconstruction and engineering knowledge','II.C. Intelligent Geometry Generation with Knowledge Augmentation, pp. 6–7; Fig. 2',[
 'Describe reconstruction from the original coordinates and connectivity, followed by material/section/support audit. Use Figure 2 to connect global geometry, structural families and a local connection.',
 'Give the equivalent material values from the INP: MAT1/MAT2 E = 120 GPa, areas 22298.692/8402.9797 mm² and their equivalent densities; MAT3 E = 206 GPa, rectangular section 98.954535 × 98.954535 mm and near-zero assigned density.',
 'Describe 8,984 initial-stress records across 1,123 cable elements, the global tensor convention and the retained reference geometry. Link numerical values to the material table and asset hashes.'
]),
('II.D. Gmsh mesh generation and identity mapping','II.D. Adaptive Mesh Generation, pp. 7–8; Fig. 3; Table II',[
 'Explain point/line creation from INP IDs, section physical groups, setTransfiniteCurve(eid, 2), first-order generate(1), and MSH 4.1 export.',
 'Distinguish geometric tags from mesh tags. Describe the explicit node/element mapping and checks: node count, element count, coordinates and ordered endpoints. Table II reports the measured mesh identity results.'
]),
('II.E. Nonlinear static analysis and load-history execution','II.E. Automated FEA Simulation and Analysis, pp. 8–9',[
 'Write equilibrium as R(u, λ) = fint(u, σ0, T) − fext(λ) = 0. Describe geometric nonlinearity, initial prestress, gravity and prescribed supports within the planar X–Z system.',
 'State the six combinations in Table I. P1 contains one dead-load step; P2–P6 independently establish dead load before their target total combination. Explain the second-step CLOAD total convention, with P5 as the worked example.',
 'Give g = 9806 mm/s², alpha = 1.2e-5 /°C for P3/P6, temperature changes −15/−34°C and the 1,125 UY constraints. Identify the locked solver build and final-time checks.'
]),
('II.F. Field recovery and automated verification','II.F. Intelligent Data Analysis and Batch Processing, p. 9',[
 'Define Umax = max sqrt(UX²+UY²+UZ²) over original nodes. Define the deformed chord n and recover N = A0 nᵀ σbar n / 1000 in kN, averaging the eight integration-point tensors.',
 'Define eight reporting span groups, their maxima and K = Nbreak/Nmax. Set Nbreak = 38080/14280 kN and give the case-specific K requirements.',
 'Define signed relative difference δ = 100(qcalc/qref − 1). Identify the independent MAPDL displacement source and the report force tables separately. Describe Figure 4 and the machine-readable evidence files.'
]),
('II.G. Reproducible batch execution and report generation','II.G. Scalable Batch Processing Architecture, pp. 9–10',[
 'Describe independent case directories, common geometry audit, recorded solver commands, failure preservation and per-stage status.',
 'Trace one final response value into CSV, VTU, PNG and the engineering report. Describe the shared result model used for Word/PDF and the style/font assets reproducing the reference report.',
 'State the deliverable structure: six solved cases, 42 field plots, comparison data, Gmsh mapping, hashes and a 17-page engineering report in both formats.'
]),
('III.A. Task specification and six-case execution','III.A. Natural Language Input and System Understanding, p. 11; Table I',[
 'Present the task request and its mapped case/output specification. Table I enumerates the six combinations and required safety factors.',
 'For the GLM session, insert the recorded model ID and stage-completion log. Report execution and retry counts from that log.'
]),
('III.B. Model reconstruction results','III.B. Geometry Generation Results, pp. 11–12; Fig. 2',[
 'Report 1,125 original nodes, 1,123 T3D2 cable elements and 71 B31 equivalent members. Show the complete geometry and member assignments in Figure 2.',
 'Report the common geometric/physical definitions across cases and the preserved initial-stress record count.'
]),
('III.C. Mesh identity and mapping results','III.C. Mesh Generation and Quality Analysis, pp. 12–13; Table II; Fig. 3',[
 'Report identical node/element counts, zero coordinate discrepancy and matching ordered connectivity between the input and the Gmsh mesh. Use the local inset to display the mapping.',
 'Explain how this identity connects solved fields to the displayed geometry and reporting groups.'
]),
('III.D. Nonlinear response and benchmark comparison','III.D. FEA Simulation and Results Analysis, p. 14; Table III',[
 'Report completion of all six cases and matching peak-displacement node IDs. Present the six numerical rows in Table III.',
 'Give maximum absolute differences: displacement 0.859887%, bottom-cable span force 0.853551% and portal-cable span force 1.125518%. Identify the controlling case/span from the supplied evidence table.',
 'Discuss P4 as the largest peak-displacement response, P5 peak at node 306, and the P3/P6 cooling responses using the solved fields.'
]),
('III.E. Cable response interpretation and safety assessment','III.E. Intelligent Result Analysis and Autonomous Design Evaluation, pp. 14–16; Fig. 4',[
 'Follow the four panels in order: global displacement agreement, span-force agreement, cable safety factors and the spatial/case distribution of signed differences.',
 'Compare the governing bottom/portal factors by case with Table I requirements. Use the heatmap to locate differences while retaining cable family and span identity.',
 'Explain the load-history correction through the P5 input audit: second-step total nodal load includes the first-step permanent contribution. Connect the input definition with the current node-306 downward response.'
]),
('III.F. Execution performance and automation accounting','III.F. Computational Performance and Scalability, pp. 16–17',[
 'Report the recorded solver wall times by case and their sum, with Python/platform/solver identity. Label timing boundaries explicitly.',
 'Add the planned GLM session measurements in a separate table: planning time, tool time, report-generation time, total session time, token usage when exposed by the provider, retries and completion count.',
 'For repeated-session evaluation, use the same six inputs and output contract; record each session independently before calculating completion rate and timing distributions.'
]),
('III.G. Engineering case presentation and report delivery','III.G. Industrial Validation: CalculiX Configuration Generation, pp. 17–18; Fig. 5',[
 'Present Figure 5 as the engineering field plate: three representative cases and two response quantities. Compare shared field scales within each row.',
 'Show how the prescribed report style is generated from the same result set, and point to the supplementary six-case contour plates, Word/PDF report and raw result archive.'
]),
('IV. Conclusions','IV. Conclusions, pp. 18–19',[
 'Conclusion 1 — Summarize the Skill-guided sequence connecting versioned bridge inputs to mesh verification, nonlinear solution and engineering-report generation.',
 'Conclusion 2 — State the measured mesh identity and six-case response differences with their actual values.',
 'Conclusion 3 — State the cable safety-factor outcomes and delivered evidence package. Add session-level GLM findings after inserting the measured session table.'
])]

def document(title, subtitle):
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.75)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6)
 for sty in ['Title','Heading 1','Heading 2','Heading 3']:
  d.styles[sty].font.name='Times New Roman';d.styles[sty].font.color.rgb=None
 d.add_heading(title,0);d.add_paragraph(subtitle)
 footer=sec.footer.paragraphs[0];footer.alignment=2
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 return d

def table(d,headers,rows):
 t=d.add_table(rows=1, cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for run in p.runs:run.font.size=Pt(9)
 return t

# Figure 4: all points are calculated from the existing validated summary.
plt.rcParams.update({'font.family':'DejaVu Serif','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
fig,ax=plt.subplots(2,2,figsize=(11.5,8.2),layout='constrained');x=np.arange(6);cases=[s['case'] for s in SUM]
ax[0,0].bar(x-.18,[s['reference_umax_mm'] for s in SUM],.36,label='MAPDL',color='#bcc6cc');ax[0,0].bar(x+.18,[s['umax_mm'] for s in SUM],.36,label='CalculiX',color='#246a88');ax[0,0].set(xticks=x,xticklabels=cases,ylabel='Maximum displacement (mm)',title='(a) Six-case displacement comparison');ax[0,0].legend(frameon=False)
for family,col,marker in [('bottom','#246a88','o'),('portal','#c06b3e','s')]:
 z=[v for s in SUM for v in s['span_forces'] if v['type']==family]
 ax[0,1].scatter([v['report_kN'] for v in z],[v['N_kN'] for v in z],s=25,c=col,marker=marker,label=family.capitalize(),alpha=.8)
ax[0,1].plot([2500,14000],[2500,14000],color='.4',ls='--',lw=.8);ax[0,1].set(xlabel='Engineering report (kN)',ylabel='Calculated span maximum (kN)',title='(b) Span-force parity (48 pairs)');ax[0,1].legend(frameon=False)
for family,col,marker in [('bottom','#246a88','o'),('portal','#c06b3e','s')]:
 ax[1,0].plot(x,[min(v['safety_factor'] for v in s['span_forces'] if v['type']==family) for s in SUM],color=col,marker=marker,label=family.capitalize())
ax[1,0].plot(x,[s['span_forces'][0]['required_factor'] for s in SUM],c='.3',ls='--',label='Required');ax[1,0].set(xticks=x,xticklabels=cases,ylabel='Minimum span safety factor',title='(c) Cable safety factors');ax[1,0].legend(frameon=False)
errs=np.array([[v['error_percent'] for v in s['span_forces']] for s in SUM]);im=ax[1,1].imshow(errs,cmap='RdBu_r',vmin=-1.2,vmax=1.2,aspect='auto')
ax[1,1].set(xticks=np.arange(8),xticklabels=['B-N','B-M','B-S','B-A','P-N','P-M','P-S','P-A'],yticks=x,yticklabels=cases,title='(d) Signed span-force differences (%)')
for i in range(6):
 for j in range(8):ax[1,1].text(j,i,f'{errs[i,j]:.2f}',ha='center',va='center',fontsize=7,color='white' if abs(errs[i,j])>.8 else 'black')
fig.colorbar(im,ax=ax[1,1],shrink=.9);fig.savefig(OUT/'figure_4_six_case_verification.png',dpi=300);fig.savefig(OUT/'figure_4_six_case_verification.pdf');plt.close(fig)
with (OUT/'figure_4_data.csv').open('w') as f:
 w=csv.writer(f);w.writerow(['case','family','span_index','calculated_kN','report_kN','difference_percent','safety_factor','required_factor'])
 for s in SUM:
  for i,v in enumerate(s['span_forces']):w.writerow([s['case'],v['type'],i%4+1,v['N_kN'],v['report_kN'],v['error_percent'],v['safety_factor'],v['required_factor']])

study=document('FeaGPT: figure-by-figure study and catwalk content mapping','Reference-led preparation for the catwalk manuscript | English working document | 6 October 2026')
study.add_heading('Reference and reading order',1);study.add_paragraph(REF);study.add_paragraph('Downloaded reference: reference/FeaGPT.pdf, 21 pages. Read the paper in its original order: Introduction; Methods A–G; Experiments and Results A–G; Conclusions. The mapping below preserves the purpose, arrangement and evidence sequence of each figure while specifying catwalk content.')
study.add_paragraph('Reference page previews reproduce the unchanged source pages. Source: Qi, Xu and Chu (2025), FeaGPT, arXiv:2510.21993v1, licensed CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/). The catwalk figure plans and quantitative Figure 4 are new adaptations.')
for n,p,title,reading,replacement,data,caption in FIGS:
 study.add_page_break();study.add_heading(f'Figure {n} — {title} (source page {p})',1)
 study.add_picture(str(ROOT/f'reference/pages/figure-{n}-page-{p:02}.png'),width=Inches(4.5))
 study.add_paragraph(f'Source page {p}: Qi, Xu and Chu, FeaGPT (2025), CC BY 4.0; unchanged page preview.')
 study.add_heading('What the source figure does',2);study.add_paragraph(reading)
 study.add_heading('Catwalk panel-by-panel replacement',2);study.add_paragraph(replacement)
 study.add_heading('Data to use',2);study.add_paragraph(data)
 study.add_heading('Draft catwalk caption',2);study.add_paragraph(caption)
study.add_page_break();study.add_heading('Table-by-table and algorithm mapping',1)
table(study,['Source item','Source function','Catwalk replacement'],[
 ['Table I, p. 11','Parameter values and configuration count','Six load combinations, step count, temperature and required K; six cases total.'],
 ['Table II, p. 13','Mesh statistics and validity checks','Node/element counts, element families, coordinate discrepancy, ordered-connectivity checks and initial-stress record count.'],
 ['Table III, p. 14','Solver performance and response summary','Six-case Umax, peak node, signed MAPDL difference, minimum cable factors and measured solver wall time.'],
 ['Algorithm 1, p. 5','Structured JSON analysis plan','Manifest-driven task schema and ordered validate/mesh/solve/recover/compare/report stages.']])
study.add_heading('Reading-to-writing sequence',1)
for title,source,items in SECTIONS:study.add_paragraph(f'{source} → {title}. {items[0]}')
study.save(ROOT/'FeaGPT_figure_by_figure_study.docx')

paper=document('Skill-guided agentic finite element analysis for nonlinear static verification of bridge catwalks','SCI manuscript framework | Adapted section and figure structure from FeaGPT | English draft outline')
paper.add_heading('Reference structure',1);paper.add_paragraph(REF);paper.add_paragraph('Working title and outline follow the reference paper’s end-to-end workflow, method modules, staged experiments and industrial demonstration. Engineering content is supplied by the catwalk Skill and the recorded six-case calculation.')
paper.add_heading('Abstract: sentence plan',1)
for t in ['Context: automated verification of prestressed catwalk cable systems requires coordinated handling of model assets, load history, solver execution and engineering reports.',
 'Method: introduce the English Skill, versioned inputs, identity-preserving Gmsh reconstruction, nonlinear CalculiX execution and automatic field/report generation.',
 'Case: identify the Zhangjinggao Bridge catwalk model with 1,125 nodes, 1,194 elements and six dead/construction/wind/temperature combinations.',
 'Results: insert 0.859887% maximum displacement difference, 0.853551% bottom-cable span-force difference and 1.125518% portal-cable span-force difference; report matching peak nodes and the calculated safety-factor outcomes.',
 'Delivery: close with the reproducible solver evidence, 42 field plots and automatically formatted Word/PDF calculation report.']:paper.add_paragraph(t)
paper.add_paragraph('Keywords: agentic finite element analysis; executable skills; bridge catwalk; nonlinear static analysis; Gmsh; CalculiX; reproducibility.')
for title,source,items in SECTIONS:
 paper.add_heading(title,1);paper.add_paragraph('Reference anchor: '+source)
 for item in items:paper.add_paragraph(item)
paper.add_page_break();paper.add_heading('Table I. Case specification',1)
loads=['Dead load','Dead + construction','Dead + construction + cooling','Dead + construction + wind','Dead + maximum gust','Dead + construction + cooling']
table(paper,['Case','Combination','Steps','ΔT / °C','Required K'],[[s['case'],loads[i],1 if i==0 else 2,[-0,-0,-15,-0,-0,-34][i],s['span_forces'][0]['required_factor']] for i,s in enumerate(SUM)])
paper.add_heading('Table II. Mesh and model identity',1)
table(paper,['Metric','Recorded value'],[['Original / Gmsh nodes','1125 / 1125'],['Original / Gmsh line elements','1194 / 1194'],['Structural element families','1123 T3D2 + 71 B31'],['Coordinate discrepancy','0 mm'],['Ordered connectivity','All 1194 elements match'],['Initial-stress records','8984'],['Reporting span groups','4 bottom + 4 portal'],['Reported-span cable elements','1113'],['Original-node planar constraints','1125 UY constraints']])
paper.add_heading('Table III. Six-case response and solver timing',1)
rows=[]
for s,e in zip(SUM,RUN['executions']):
 rows.append([s['case'],f"{s['umax_mm']:.3f}",s['peak_node'],f"{s['displacement_error_percent']:.4f}",f"{min(v['safety_factor'] for v in s['span_forces'] if v['type']=='bottom'):.4f}",f"{min(v['safety_factor'] for v in s['span_forces'] if v['type']=='portal'):.4f}",f"{e['elapsed_s']:.3f}"])
table(paper,['Case','Umax/mm','Node','δU/%','K bottom','K portal','Solver/s'],rows)
paper.add_paragraph(f"Recorded solver wall time summed across six sequential cases: {sum(e['elapsed_s'] for e in RUN['executions']):.6f} s. Timing source: examples/run.json, executions[].elapsed_s; timing covers the solver subprocess. Run ID: {RUN['run_id']}.")
paper.add_heading('Algorithm 1. Skill-guided execution',1)
algorithm=['Read task and resolve the versioned six-case manifest.','Verify asset hashes and audit material, section, prestress, load and support cards.','Regenerate the Gmsh line mesh and verify original-to-mesh identity.','For each case P1–P6: copy its complete INP into a new case directory and execute the recorded solver command.','Validate final step/time and complete finite U/S fields.','Recover deformed-chord cable forces; evaluate span maxima, safety factors and benchmark differences.','Write CSV, VTU and seven field plots per case.','Generate the Chinese engineering report from the common result object and prescribed style assets.','Write execution status and output hashes; return the report and evidence links.']
for i,t in enumerate(algorithm,1):paper.add_paragraph(f'{i}. {t}')
paper.add_heading('Planned GLM session table',1)
table(paper,['Field','Value source'],[['Provider/model/version','Actual GLM session metadata'],['Request and tool trace','Saved request and tool-call log'],['Planning / tools / report / total time','Session stage timestamps'],['Token usage','Provider usage record'],['Retries and completed stages','Execution log'],['Repeated-session completion rate','Completed sessions / attempted sessions']])
paper.add_heading('Figure list and captions',1)
for n,p,title,reading,replacement,data,caption in FIGS:
 paper.add_heading(f'Figure {n}',2);paper.add_paragraph(caption);paper.add_paragraph('Panel construction: '+replacement)
 if n==4:paper.add_picture(str(OUT/'figure_4_six_case_verification.png'),width=Inches(6.6));paper.add_paragraph('B/P = bottom/portal cables; N/M/S/A = north side/main/south side/south auxiliary span.')
paper.add_heading('Supplementary material',1)
for t in ['S1: English Skill, seven execution references, source scripts and environment requirements.','S2: Input manifest, asset hashes, material/load audits and Gmsh mapping.','S3: Complete 48-row span-force comparison and six-case displacement table.','S4: Six-case solver logs, DAT/FRD/STA, CSV/VTU and 42 field plots.','S5: Generated Word/PDF engineering report and typography specification.','S6: GLM session trace and repeated-session measurements, populated from the planned execution experiment.']:paper.add_paragraph(t)
paper.add_heading('References and data provenance',1);paper.add_paragraph(REF);paper.add_paragraph('Zhangjinggao Yangtze River Bridge, South Navigation Channel Bridge: Catwalk Structural Review Calculation Report, version 0324, 118 pages. Project-supplied PDF; force benchmarks from Tables 1-11 and 1-14, safety-factor criteria from Table 1-7.');paper.add_paragraph('Catwalk static review Skill v1.2.0. Numerical evidence: examples/summary.json, examples/comparison.csv and examples/run.json. Input, solver and PDF identities are recorded in the package manifests.')
paper.save(ROOT/'SCI_manuscript_framework.docx')

# Editable LaTeX outline with the same reference-led paragraph plan.
def esc(s):
 for a,b in [('\\',r'\textbackslash{}'),('&',r'\&'),('%',r'\%'),('_',r'\_'),('#',r'\#')]:s=s.replace(a,b)
 return s
tex=[r'\documentclass[11pt]{article}',r'\usepackage{fontspec}',r'\usepackage[margin=25mm]{geometry}',r'\usepackage{graphicx}',r'\usepackage{hyperref}',r'\title{Skill-guided agentic finite element analysis for nonlinear static verification of bridge catwalks}',r'\date{}',r'\begin{document}',r'\maketitle',r'\noindent Reference-led manuscript outline. Compile with XeLaTeX or LuaLaTeX.']
for title,source,items in SECTIONS:
 tex += [r'\section*{'+esc(title)+'}',r'\textit{Reference anchor: '+esc(source)+'}',r'\begin{enumerate}']+[r'\item '+esc(item) for item in items]+[r'\end{enumerate}']
tex += [r'\begin{figure}[ht]',r'\centering',r'\includegraphics[width=\textwidth]{figures/figure_4_six_case_verification.pdf}',r'\caption{'+esc(FIGS[3][-1])+'}',r'\end{figure}',r'\begin{thebibliography}{9}',r'\bibitem{qi2025feagpt} '+esc(REF),r'\end{thebibliography}',r'\end{document}']
(ROOT/'manuscript_outline.tex').write_text('\n'.join(tex)+'\n')
with (ROOT/'figure_mapping.csv').open('w') as f:
 w=csv.writer(f);w.writerow(['figure','source_page','purpose','source_reading','catwalk_replacement','data','draft_caption']);w.writerows(FIGS)
(ROOT/'section_mapping.json').write_text(json.dumps([dict(section=t,reference=s,paragraph_plan=i) for t,s,i in SECTIONS],indent=2,ensure_ascii=False)+'\n')
(ROOT/'task_specification.json').write_text(json.dumps({'skill':'catwalk-static-review','skill_version':'1.2.0','cases':cases,'units':{'length':'mm','force':'N','stress':'MPa','mass':'tonne','temperature':'degC'},'input_manifest':'../../skills/catwalk-static-review/assets/reference/manifest.json','solver_sha256':RUN['solver_sha256'],'stages':algorithm,'outputs':['DOCX','PDF','CSV','VTU','PNG','DAT','FRD','STA','JSON','SHA256SUMS']},indent=2)+'\n')
maxima={}
for fam in ['bottom','portal']:
 pairs=[(s['case'],v) for s in SUM for v in s['span_forces'] if v['type']==fam];case,v=max(pairs,key=lambda z:abs(z[1]['error_percent']));maxima[fam]={'case':case,**v}
(ROOT/'evidence_map.json').write_text(json.dumps({'run_id':RUN['run_id'],'summary_sha256':hashlib.sha256((SKILL/'examples/summary.json').read_bytes()).hexdigest(),'max_abs_displacement_difference_percent':max(abs(s['displacement_error_percent']) for s in SUM),'governing_force_differences':maxima,'solver_wall_time_sum_s':sum(e['elapsed_s'] for e in RUN['executions']),'figure_4_source':'../../skills/catwalk-static-review/examples/summary.json','session_experiment':{'provider_preference':'GLM','status':'planned','fields':'See manuscript framework session table'}},ensure_ascii=False,indent=2)+'\n')
(ROOT/'references.bib').write_text('''@misc{qi2025feagpt,
  title={FeaGPT: an End-to-End agentic-AI for Finite Element Analysis},
  author={Qi, Yupeng and Xu, Ran and Chu, Xu},
  year={2025},
  eprint={2510.21993},
  archivePrefix={arXiv},
  url={https://arxiv.org/abs/2510.21993v1}
}
''')
print('Built manuscript framework, figure-by-figure study, LaTeX outline, Figure 4 and evidence mappings.')

# Portable PDF reading copies generated from the same DOCX content.
from io import BytesIO
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from docx.text.paragraph import Paragraph as WordParagraph
from docx.table import Table as WordTable
styles=getSampleStyleSheet()
for name in ['Normal','BodyText','Title','Heading1','Heading2','Heading3']:
 styles[name].fontName='Times-Roman' if name in ['Normal','BodyText'] else 'Times-Bold'
styles['Normal'].fontSize=10;styles['Normal'].leading=13
styles['Normal'].spaceAfter=6
styles['Heading1'].fontSize=14;styles['Heading1'].leading=17

def pdfcopy(src):
 doc=Document(src);story=[]
 for block in doc.element.body:
  if block.tag==qn('w:p'):
   p=WordParagraph(block,doc)
   for blip in block.findall('.//'+qn('a:blip')):
    data=doc.part.related_parts[blip.get(qn('r:embed'))].blob
    extent=block.find('.//'+qn('wp:extent'))
    w=float(extent.get('cx'))/12700;h=float(extent.get('cy'))/12700
    scale=min(1,465/w,520/h);story.append(Image(BytesIO(data),width=w*scale,height=h*scale))
   if p.text:
    name={'Title':'Title','Heading 1':'Heading1','Heading 2':'Heading2','Heading 3':'Heading3'}.get(p.style.name,'Normal')
    text=p.text.replace('→',' > ').replace('−','-').replace('σ','sigma').replace('λ','lambda').replace('Δ','Delta ').replace('δ','delta ').replace('ᵀ','^T').replace('⊗',' tensor ')
    story.append(Paragraph(escape(text),styles[name]))
   if any(e.get(qn('w:type'))=='page' for e in block.findall('.//'+qn('w:br'))):story.append(PageBreak())
  elif block.tag==qn('w:tbl'):
   t=WordTable(block,doc);rows=[[Paragraph(escape(c.text).replace('Δ','Delta ').replace('δ','delta '),styles['Normal']) for c in row.cells] for row in t.rows]
   tb=Table(rows,colWidths=[465/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
   tb.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.35,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef0f1')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]));story.extend([tb,Spacer(1,9)])
 def footer(canvas,doc):
  canvas.setFont('Times-Roman',9);canvas.drawRightString(530,28,str(doc.page))
 SimpleDocTemplate(str(src.with_suffix('.pdf')),pagesize=(595.28,841.89),leftMargin=65,rightMargin=65,topMargin=48,bottomMargin=45,title=src.stem).build(story,onFirstPage=footer,onLaterPages=footer)
for name in ['SCI_manuscript_framework.docx','FeaGPT_figure_by_figure_study.docx']:pdfcopy(ROOT/name)
print('Created both PDF reading copies.')
